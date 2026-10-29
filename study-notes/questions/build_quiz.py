import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_BREAK
from q_data import Q

def base(title):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    for m in ("left_margin","right_margin","top_margin","bottom_margin"):
        setattr(s, m, Cm(1.9))
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    d.add_heading(title, 0)
    return d

def runs(p, text, bold=False):
    for i, part in enumerate(re.split(r"`", text)):
        r = p.add_run(part); r.bold = bold
        if i % 2: r.font.name = "Consolas"

def tight(p, after=2):
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(0)

d = base("CCNA Practice Questions: Set 1 (Q1 to Q100)")
p = d.add_paragraph("Answers are in the separate file CCNA_Set1_Answer_Key.docx. Mark your answer, then check.")
p.runs[0].italic = True
for n, (q, o, a, w, c) in enumerate(Q, 1):
    p = d.add_paragraph(); runs(p, f"Q{n}. {q}", True); tight(p, 1); p.paragraph_format.space_before = Pt(8); p.paragraph_format.keep_with_next = True
    for L, t in zip("ABCD", o):
        p = d.add_paragraph(); runs(p, f"    {L}.  {t}"); tight(p, 0); p.paragraph_format.keep_with_next = L != "D"
d.save("CCNA_Set1_Questions.docx")

k = base("CCNA Practice Questions: Set 1 Answer Key")
p = k.add_paragraph("Do not open until you have answered. Items marked CHECK may be a flawed question or a garbled transcript. Verify them against Jeremy's videos before trusting the answer.")
p.runs[0].italic = True
t = k.add_table(rows=1, cols=4); t.style = "Table Grid"
for c, h in zip(t.rows[0].cells, ["Q", "Ans", "Why", "CHECK"]):
    c.text = ""; c.paragraphs[0].add_run(h).bold = True
for n, (q, o, a, w, c) in enumerate(Q, 1):
    r = t.add_row().cells
    r[0].text = str(n); r[1].text = a; r[2].text = w.replace("`", ""); r[3].text = c
    if c: r[3].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
for row in t.rows:
    for c, wd in zip(row.cells, [1.1, 1.2, 7.0, 7.9]):
        c.width = Cm(wd)
        for p in c.paragraphs: tight(p, 0); [setattr(r.font, "size", Pt(9)) for r in p.runs]
k.save("CCNA_Set1_Answer_Key.docx")
print(sum(1 for x in Q if x[4]), "flagged")
